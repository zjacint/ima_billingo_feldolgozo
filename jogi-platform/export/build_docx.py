# Kérdéslista Word-dokumentum. Futtatás: python3 build_docx.py
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT

Q = [
("Üzlet és jog", [
 ("Ki a jogi szolgáltató (egyéni ügyvéd / iroda)? Hány ügyvéd és asszisztens használja majd az admin felületet?", "Meghatározza a szerepköröket és a jogosultsági modellt."),
 ("Végrendelet: sablont szeretnétek, vagy csak előkészítő kérdőívet ügyvédi véglegesítéssel? (A formai szabályok szigorúak: kézírás, két tanú, ellenjegyzés.)", "Eldönti, kell-e „ügyvédi ellenjegyzés szükséges” folyamat a sablonokhoz."),
 ("Mely szerződéstípusok kerüljenek az induló katalógusba (kb. 10–15)? Pl. adásvétel, bérlet, megbízás, vállalkozás, munkaszerződés, titoktartás, tartozáselismerés, ajándékozás.", "Ebből készül az első sablonkészlet."),
 ("Mit tartalmaz pontosan az előfizetés? Korlátlan sablonhasználat vagy havi keret? Benne van-e ügyvédi konzultáció?", "Csomag- és keretmodell (Plan, PlanEntitlement)."),
 ("Árazás: havi/éves díj, külön csomag magánszemélynek és cégnek? Van tervezett ársáv?", ""),
 ("Szükséges-e tőletek ÁSZF, „nem minősül jogi tanácsnak” kizáró szöveg, felelősségbiztosítási információ? Van kész szövegetek?", ""),
 ("Egyeztettetek-e a Kamarával az online jogi szolgáltatás és hirdetés szabályairól?", "Megfelelőségi kockázat."),
]),
("Felhasználók és adatok", [
 ("Célcsoport: magánszemélyek, cégek, vagy mindkettő? Kell-e több felhasználó egy céghez?", "Organization/OrgMember modell."),
 ("Milyen személyes adatok kellenek a szerződésekhez (adóazonosító, anyja neve, lakcím, személyi igazolvány száma)?", "Adatminimalizálás, mező-szintű titkosítás."),
 ("Mennyi ideig kell megőrizni az ügyfél-dokumentumokat és az ügyeket? Mikor töröljünk?", "File.deleteAfter, törlési folyamat."),
 ("Kell-e elektronikus aláírás (AVDH, KAÜ, DocuSign / Adobe Sign), vagy elég a PDF-letöltés?", "Integrációs terv."),
 ("Kell-e NAV- vagy ügyfélkapu-integráció (adózói adatok betöltése), vagy elég a tudástár?", ""),
]),
("Tudástár és AI", [
 ("Ki írja és tartja naprakészen a tudástár tartalmát? Kell-e automatikus jogszabályfigyelés (njt.hu, NAV)?", "Szinkron-folyamat és felülvizsgálati határidők."),
 ("A chat szabad szöveggel válaszoljon a tudástár alapján, vagy csak előre megírt GYIK-ből?", "Kockázat és költség."),
 ("A Gemini-felülvizsgálat csak a saját sablonjainkra fusson, vagy az ügyfél által feltöltött idegen szerződésre is?", "Az utóbbi nagyobb kockázat és ügyvédi ellenőrzési igény."),
 ("Minden AI-riportot ügyvéd néz át, vagy megjelenhet azonnal figyelmeztetéssel?", "Ügyvédi ellenőrzési sor mérete."),
 ("Elfogadható-e a Google Vertex AI (EU régió) használata az ügyvédi titoktartás szempontjából? Kell-e ügyfél-hozzájárulás az AI-feldolgozáshoz?", "Adatkezelési és titoktartási döntés."),
]),
("Fizetés és számlázás", [
 ("Fizetési szolgáltató: Stripe, Barion vagy OTP SimplePay?", ""),
 ("Számlázás: Billingo, Számlázz.hu vagy más?", ""),
 ("Az árajánlat elfogadása után online fizetés kell, vagy átutalás is megfelel?", ""),
]),
("Technikai és arculati", [
 ("Csak magyar nyelv, vagy később angol is?", ""),
 ("Van meglévő GCP-szervezet/projekt? Régiópreferencia (javaslat: europe-west3 Frankfurt vagy europe-central2 Varsó)?", ""),
 ("Van márkanév, logó, színvilág, saját domain?", ""),
 ("Rendben van a Next.js + Prisma + PostgreSQL + Cloud Run technológiai javaslat?", ""),
 ("A mintaoldalak (NCOA, CNBC, LawDepot) melyik elemei legyenek kötelezően megvalósítva (pl. lépésenkénti kérdőív, összehasonlító táblázat)?", ""),
 ("Mikorra tervezitek a pilotot, és hány pilot ügyfél van?", "Ütemezés."),
]),
]

d = Document()
s = d.sections[0]; s.left_margin = s.right_margin = Cm(2); s.top_margin = s.bottom_margin = Cm(2)
st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
t = d.add_heading("Ügyvédi előfizetéses platform – nyitott kérdések", 0)
d.add_paragraph("Kérjük, a „Válasz” oszlopba írjatok röviden. Ahol nem tudtok még dönteni, jelöljétek „később”-nek – ezzel is tudunk haladni. A „Miért kérdezzük?” oszlop megmutatja, mire hat a válasz.")
n = 1
for sec, items in Q:
    d.add_heading(sec, 1)
    tb = d.add_table(rows=1, cols=4); tb.style = "Light Grid Accent 1"; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, h in zip(tb.rows[0].cells, ["#", "Kérdés", "Miért kérdezzük?", "Válasz"]):
        c.text = h
    for q, why in items:
        r = tb.add_row().cells
        r[0].text = str(n); r[1].text = q; r[2].text = why; r[3].text = ""
        n += 1
    for row in tb.rows:
        for c, w in zip(row.cells, [Cm(0.9), Cm(7.2), Cm(4), Cm(5)]):
            c.width = w
d.save("ugyvedi_platform_kerdesek.docx")
